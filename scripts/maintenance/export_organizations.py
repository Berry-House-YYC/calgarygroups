#!/usr/bin/env python3
"""Export all organization listings to CSV and formatted XLSX.

Reads every markdown file under src/content/organizations/, parses the YAML
front matter (without requiring PyYAML, since the front matter is
machine-generated and simple), extracts contact details and notes from the
body, and writes:

  docs/calgary-groups.csv
  docs/calgary-groups.xlsx

Usage:
    python3 scripts/maintenance/export_organizations.py
"""

import csv
import glob
import html
import os
import re
import sys

ORGS_DIR = os.path.join('src', 'content', 'organizations')
DOCS_DIR = 'docs'

COLUMNS = [
    'Name',
    'Type',
    'Status',
    'Interests',
    'Age Range',
    'Identity Focused',
    'Meeting Format',
    'Location Area',
    'Description',
    'Website',
    'Email',
    'Phone',
    'Address',
    'Hours',
    'Notes',
    'Additional Info',
    'Community Submitted',
    'Folder',
    'Slug',
    'Site Path',
]

CONTACT_LABELS = {'website', 'email', 'phone', 'address', 'hours'}


def parse_scalar(value):
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in ('"', "'"):
        value = value[1:-1]
    low = value.lower()
    if low == 'true':
        return True
    if low == 'false':
        return False
    return value


def parse_frontmatter(text):
    """Parse the simple front matter used by these files.

    Handles `key: value`, `key: ["a", "b"]`, and folded (indented
    continuation) scalars. Returns (dict, body).
    """
    m = re.match(r'^---\s*\n(.*?)\n---\s*\n?', text, re.S)
    if not m:
        return {}, text

    data = {}
    last_key = None
    for line in m.group(1).split('\n'):
        if not line.strip():
            continue
        km = re.match(r'^([A-Za-z_][A-Za-z_ ]*):\s*(.*)$', line)
        if km and not line[0].isspace():
            key = km.group(1).strip().replace(' ', '_')
            raw = km.group(2).strip()
            if raw.startswith('[') and raw.endswith(']'):
                items = [parse_scalar(i) for i in raw[1:-1].split(',') if i.strip()]
                data[key] = [str(i) for i in items]
            else:
                data[key] = parse_scalar(raw)
            last_key = key
        elif last_key is not None:
            # folded continuation line of a multi-line scalar
            prev = data[last_key]
            cont = line.strip()
            data[last_key] = (str(prev) + ' ' + cont).strip() if prev else cont

    return data, text[m.end():]


def unescape_entities(text):
    """Unescape only entities that end with ';' so URLs like
    '...&ltsid=...' (a real Linktree param) are not mangled by
    html.unescape's legacy no-semicolon entity handling."""
    return re.sub(r'&(?:#\d+|#x[0-9a-fA-F]+|[a-zA-Z]+);',
                  lambda m: html.unescape(m.group(0)), text)


def strip_tags(fragment):
    fragment = re.sub(r'<[^>]+>', '', fragment)
    return unescape_entities(fragment).strip()


def strip_markdown_links(text):
    # [text](url) -> text ; [(403) 608-2401](<>) -> (403) 608-2401
    return re.sub(r'\[([^\]]*)\]\([^)]*\)', r'\1', text)


def parse_body(body):
    contact = {label: '' for label in CONTACT_LABELS}
    rest = body

    # HTML contact block: <li>Website: <a href="...">...</a></li>
    for li in re.findall(r'<li>(.*?)</li>', rest, re.S):
        lm = re.match(r'\s*([A-Za-z ]+):\s*(.*)', li, re.S)
        if not lm:
            continue
        label = lm.group(1).strip().lower()
        if label in CONTACT_LABELS:
            contact[label] = strip_tags(lm.group(2))

    # Remove the whole contact div so it doesn't leak into notes
    rest = re.sub(r'<div class="org-contact-info">.*?</div>', '', rest, flags=re.S)

    # Markdown-style contact block: **Contact Info:** then "- Email: x"
    md_contact = re.search(r'\*\*Contact Info:\*\*(.*?)(?=\n\*\*|\Z)', rest, re.S)
    if md_contact:
        for li in re.findall(r'-\s*([A-Za-z ]+):\s*(.*)', md_contact.group(1)):
            label = li[0].strip().lower()
            if label in CONTACT_LABELS and not contact[label]:
                contact[label] = strip_markdown_links(li[1]).strip()
        rest = rest.replace(md_contact.group(0), '')

    # **Notes:** section (usually last) — markdown and HTML variants
    notes = ''
    notes_m = re.search(r'\*\*Notes:\*\*\s*(.*)', rest, re.S)
    if notes_m:
        notes = strip_markdown_links(notes_m.group(1)).strip()
        rest = rest[:notes_m.start()]

    html_notes = re.search(
        r'<div class="org-notes">\s*<strong>Notes:</strong>(.*?)</div>', rest, re.S)
    if html_notes:
        notes = (notes + '\n' if notes else '') + strip_tags(html_notes.group(1))
        rest = rest.replace(html_notes.group(0), '')

    # Whatever remains is description + occasional extra paragraphs
    extra = rest.strip()
    return contact, notes, extra


def normalize_ws(text):
    return re.sub(r'\s+', ' ', text or '').strip()


def collect_rows():
    rows = []
    skipped = []
    for path in sorted(glob.glob(os.path.join(ORGS_DIR, '*', '*.md'))):
        if os.path.getsize(path) == 0:
            skipped.append(path)
            continue
        with open(path, encoding='utf-8') as f:
            text = f.read()

        fm, body = parse_frontmatter(text)
        # normalize the one-off "community submission" key variant
        if 'community_submission' in fm and 'community_submitted' not in fm:
            fm['community_submitted'] = fm.pop('community_submission')
        contact, notes, extra = parse_body(body)

        # Drop the leading body paragraph if it just repeats the description
        desc = str(fm.get('description', '') or '')
        paragraphs = [p.strip() for p in extra.split('\n\n') if p.strip()]
        if paragraphs and normalize_ws(strip_markdown_links(paragraphs[0])) == normalize_ws(desc):
            paragraphs = paragraphs[1:]
        additional = '\n\n'.join(strip_markdown_links(p) for p in paragraphs)

        slug = os.path.splitext(os.path.basename(path))[0]
        folder = os.path.basename(os.path.dirname(path))
        interests = fm.get('interests', [])
        if isinstance(interests, str):
            interests = [interests]

        rows.append({
            'Name': fm.get('name', ''),
            'Type': fm.get('type', ''),
            'Status': fm.get('status', ''),
            'Interests': '; '.join(interests),
            'Age Range': fm.get('age_range', ''),
            'Identity Focused': 'Yes' if fm.get('identity_focused') is True else (
                'No' if fm.get('identity_focused') is False else ''),
            'Meeting Format': fm.get('meeting_format', ''),
            'Location Area': fm.get('location_area', ''),
            'Description': desc,
            'Website': contact['website'],
            'Email': contact['email'],
            'Phone': contact['phone'],
            'Address': contact['address'],
            'Hours': contact['hours'],
            'Notes': notes,
            'Additional Info': additional,
            'Community Submitted': 'Yes' if fm.get('community_submitted') else '',
            'Folder': folder,
            'Slug': slug,
            'Site Path': f'/organizations/{slug}/',
        })
    return rows, skipped


def write_csv(rows, path):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Organizations'

    header_fill = PatternFill('solid', fgColor='2B2B2B')
    header_font = Font(bold=True, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='top')
    top = Alignment(vertical='top')

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 22

    wide_cols = {'Description', 'Notes', 'Additional Info'}
    for row in rows:
        ws.append([row[c] for c in COLUMNS])

    for idx, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        if col in wide_cols:
            ws.column_dimensions[letter].width = 50
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).alignment = wrap
        else:
            longest = max(
                [len(str(ws.cell(row=r, column=idx).value or '')) for r in range(1, ws.max_row + 1)]
                or [0]
            )
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), 40)
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).alignment = top

    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    summary = wb.create_sheet('Summary')
    summary.append(['Field', 'Value', 'Count'])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font

    def counts(field):
        tally = {}
        for row in rows:
            v = row[field] or '(blank)'
            tally[v] = tally.get(v, 0) + 1
        return sorted(tally.items(), key=lambda kv: (-kv[1], kv[0]))

    interest_tally = {}
    for row in rows:
        for i in (row['Interests'].split('; ') if row['Interests'] else []):
            interest_tally[i] = interest_tally.get(i, 0) + 1

    for field in ('Type', 'Status', 'Age Range', 'Meeting Format', 'Location Area'):
        for value, count in counts(field):
            summary.append([field, value, count])
    for value, count in sorted(interest_tally.items(), key=lambda kv: (-kv[1], kv[0])):
        summary.append(['Interests', value, count])

    summary.column_dimensions['A'].width = 18
    summary.column_dimensions['B'].width = 30
    summary.column_dimensions['C'].width = 8
    summary.auto_filter.ref = summary.dimensions

    wb.save(path)


def main():
    rows, skipped = collect_rows()
    os.makedirs(DOCS_DIR, exist_ok=True)
    csv_path = os.path.join(DOCS_DIR, 'calgary-groups.csv')
    xlsx_path = os.path.join(DOCS_DIR, 'calgary-groups.xlsx')
    write_csv(rows, csv_path)

    try:
        write_xlsx(rows, xlsx_path)
    except ImportError:
        print('openpyxl not installed; wrote CSV only. pip install openpyxl for XLSX.')
        xlsx_path = None

    print(f'Exported {len(rows)} organizations')
    print(f'  CSV:  {csv_path}')
    if xlsx_path:
        print(f'  XLSX: {xlsx_path}')
    for path in skipped:
        print(f'  skipped empty file: {path}')


if __name__ == '__main__':
    sys.exit(main())
